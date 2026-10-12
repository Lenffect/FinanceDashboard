#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Eenmalige import van Factuur.xlsx, Finance 2026.xlsx en Finance.xlsx in de administratie-database.
Gebruik (vereist openpyxl):  python3 importeer_spreadsheets.py <map met de xlsx-bestanden> --reset
LET OP: --reset maakt de database leeg en begint opnieuw."""
import sys, os, re, datetime as dt
from contextlib import closing
import openpyxl

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, os.path.dirname(HERE))
import server as S  # noqa

MAANDEN = ['januari', 'februari', 'maart', 'april', 'mei', 'juni', 'juli', 'augustus', 'september', 'oktober', 'november', 'december']


def iso(v):
    if isinstance(v, dt.datetime):
        return v.date().isoformat()
    if isinstance(v, dt.date):
        return v.isoformat()
    s = str(v).strip().replace('Sept', 'Sep')
    return dt.datetime.strptime(s, '%d %b %Y').date().isoformat()


def periode_tekst(v):
    if isinstance(v, (dt.datetime, dt.date)):
        return '%s %d' % (MAANDEN[v.month - 1].capitalize(), v.year)
    return str(v or '')


def kv(s):  # "KvK nr: 123" -> "123"
    return re.sub(r'^[^:]*:\s*', '', str(s or '')).strip()


def main(folder, reset):
    if reset and os.path.exists(S.DB):
        os.remove(S.DB)
    S.init_db()
    fac = openpyxl.load_workbook(os.path.join(folder, 'Factuur.xlsx'))
    f26 = openpyxl.load_workbook(os.path.join(folder, 'Finance 2026.xlsx'))['Sheet1']
    fin = openpyxl.load_workbook(os.path.join(folder, 'Finance.xlsx'))
    exp, wise = fin['Expenses'], fin['Wise card Cash Flow']
    c = S.db()

    def boeking(**kw):
        b = S.clean_boeking(kw)
        return S.insert(c, 'boekingen', S.BOEK_FIELDS, b)

    # ---- relaties + facturen aan klanten (FLUID-bladen)
    betaald_op = {'202601': '2026-04-11', '202602': '2026-05-10', '202603': '2026-06-13', '202604': '2026-06-11',
                  '202605': '2026-07-13', '202606': '2026-07-16', '202607': '2026-08-11', '202609': '2026-10-07'}
    rel_ids = {}
    sheets = sorted([ws for ws in fac if ws.title.startswith('(FLUID)')], key=lambda w: str(int(w['G10'].value)))
    for ws in sheets:
        naam = ws['A6'].value
        klant = {'naam': naam, 'adres': ws['A7'].value or '', 'postcode_plaats': ws['A8'].value or '',
                 'kvk': kv(ws['A9'].value), 'btw_nr': kv(ws['A10'].value), 'email': ws['B4'].value or ''}
        if naam not in rel_ids:
            rel_ids[naam] = S.insert(c, 'relaties', S.REL_FIELDS, {**klant, 'soort': 'klant', 'std_btw_soort': 'nl21', 'notitie': ''})
        else:  # nieuwste e-mailadres bewaren
            c.execute('UPDATE relaties SET email=? WHERE id=?', (klant['email'], rel_ids[naam]))
    c.commit()
    c.close()

    for ws in sheets:
        nr = str(int(ws['G10'].value))
        naam = ws['A6'].value
        klant = {'naam': naam, 'adres': ws['A7'].value or '', 'postcode_plaats': ws['A8'].value or '',
                 'kvk': kv(ws['A9'].value), 'btw_nr': kv(ws['A10'].value), 'email': ws['B4'].value or ''}
        regels = []
        kap = ws['A15'].value == 'Werkzaamheden'
        for r in range(16, 21):
            if ws['G%d' % r].value in (None, ''):
                continue
            if kap:
                regels.append({'periode': '', 'omschrijving': ws['A%d' % r].value, 'aantal': ws['E%d' % r].value,
                               'prijs_tekst': str(ws['F%d' % r].value or ''), 'totaal': ws['G%d' % r].value})
            else:
                regels.append({'periode': periode_tekst(ws['A%d' % r].value),
                               'omschrijving': ws['E%d' % r].value or ws['D%d' % r].value, 'totaal': ws['G%d' % r].value})
        notitie = (ws['A21'].value or '') if isinstance(ws['A21'].value, str) and ws['A21'].value.startswith('*') else ''
        stamps = 'Stamps' in naam
        d = {'nummer': nr, 'datum': iso(ws['G9'].value), 'klant': klant, 'onder': 'fluid', 'regels': regels, 'btw_pct': 21,
             'termijn': 14, 'notitie': notitie, 'toewijzing': 'gedeeld' if stamps else 'ik', 'pct_ik': 50 if stamps else 100}
        rid = S.q('SELECT id FROM relaties WHERE naam=?', (naam,), one=True)['id']
        d['relatie_id'] = rid
        fid = S.save_factuur(d)
        with closing(S.db()) as cc:
            if nr in betaald_op:
                cc.execute('UPDATE boekingen SET betaald=1, betaaldatum=?, categorie=? WHERE factuur_id=?', (betaald_op[nr], 'Omzet', fid))
            else:
                cc.execute('UPDATE boekingen SET controleren=1, notitie=?, categorie=? WHERE factuur_id=?',
                           ('Staat niet als ontvangen in "Finance 2026". Betaald? Zo ja: vink betaald aan.', 'Omzet', fid))
            cc.commit()

    c = S.db()

    # ---- JPPW-facturen (Jelmer -> Lenffect)
    jppw_betaald = {'202601': '2026-04-22', '202602': '2026-05-18', '202603': '2026-06-19', '202604': '2026-08-16', '202605': None}
    for ws in fac:
        if 'JJPW' in ws.title:
            nr = str(int(ws['G10'].value))
            datum = iso(ws['G9'].value)
            excl = float(ws['G16'].value)
            notitie = 'Factuur van %s aan Lenffect (%s).' % (S.settings()['partner_naam'], ws['E16'].value)
            controleren = 0
            betaald = jppw_betaald.get(nr)
            if not betaald:
                controleren = 1
                notitie += ' Nog niet teruggevonden als betaald in je overzichten (periode 01-08 t/m 01-10). Is dit al betaald? Vink dan "betaald" aan.'
            boeking(soort='kosten', datum=datum, omschrijving='JPPW %s - %s' % (nr, ws['E16'].value), relatie=S.settings()['partner_naam'],
                    factuurnr='JPPW ' + nr, bedrag_excl=excl, btw_bedrag=round(excl * 0.21, 2), btw_soort='nl21', toewijzing='ik',
                    betaald=1 if betaald else 0, betaaldatum=betaald, rekening='ing', categorie=S.JPPW,
                    notitie=notitie, controleren=controleren)

    # ---- overige zakelijke kosten uit "Finance 2026"
    def f26_rows():
        for row in f26.iter_rows(min_row=12, max_row=39):
            b, d, e, f, g = row[1].value, row[3].value, row[4].value, row[5].value, row[6].value
            if b and d and e is not None and f in ('Uitgaand', 'Inkomend'):
                yield b, iso(d), float(e), f, g

    overig = {'KvK inschrijving': ('Oprichting', 'geen', 0), 'Gumroad': ('Software & content', 'nl21', 1), 'Embrosa': ('Overig', 'nl21', 0),
              'Adobe CC': ('Software & abonnementen', 'nl21', 1), 'Mijn host': ('Hosting & domeinen', 'nl21', 0)}
    overig = {k.lower(): v for k, v in overig.items()}
    for b, d, e, f, g in f26_rows():
        b = b.strip()
        if b.lower() in overig:
            cat, soort, ctrl = overig[b.lower()]
            btw = float(g or 0)
            note = ''
            if ctrl:
                note = 'Buitenlandse leverancier met btw: controleer op de factuur of dit Nederlandse btw is (dan aftrekbaar), of vul je btw-nummer bij de leverancier in zodat de btw wordt verlegd.'
            boeking(soort='kosten', datum=d, omschrijving=b, relatie=b, bedrag_excl=round(e - btw, 2), btw_bedrag=btw,
                    btw_soort=soort, toewijzing='ik', betaald=1, rekening='ing', categorie=cat, controleren=ctrl, notitie=note)
        elif b.startswith('BTW Q'):
            boeking(soort='btw_betaling', datum=d, omschrijving='Btw-betaling ' + b[4:], relatie='Belastingdienst', bedrag_excl=e, rekening='ing', categorie='Btw')
    boeking(soort='storting', datum='2026-03-01', omschrijving='Startsaldo zakelijke rekening (privé storting)', relatie='Privé (030)', bedrag_excl=200, rekening='ing')

    # ---- abonnementen + kosten van de Wise-kaart (sheet "Expenses")
    ab = {}

    def abo(naam, **kw):
        base = dict(naam=naam, omschrijving='', bedrag_excl=0, btw_bedrag=0, btw_soort='geen', toewijzing='ik', pct_ik=100, rekening='wise',
                    categorie='Software & abonnementen', interval='maandelijks', startdatum=None, laatste_datum=None, actief=1, variabel=0, notitie='')
        base.update(kw)
        ab[kw.get('key', naam)] = S.insert(c, 'abonnementen', S.ABO_FIELDS, base)

    abo('Google', omschrijving='Google-abonnement (maandelijks)', bedrag_excl=16.2, btw_soort='eu_verlegd', startdatum='2026-04-01',
        laatste_datum='2026-10-01', variabel=1, notitie='Google Ierland factureert zonder btw (EU, verlegd).', key='google_m')
    abo('Google', omschrijving='Domein fluid-marketing.com (jaarlijks)', bedrag_excl=12.0, btw_soort='eu_verlegd', interval='jaarlijks',
        startdatum='2026-02-08', laatste_datum='2026-02-08', categorie='Hosting & domeinen', key='google_j')
    abo('Freepik', omschrijving='Pro+', bedrag_excl=36.0, btw_bedrag=7.56, btw_soort='nl21', startdatum='2026-03-23', laatste_datum='2026-05-23',
        actief=0, notitie='Laatste betaling mei. Zet weer aan als je nog betaalt.', key='freepik')
    abo('Artlist', bedrag_excl=12.97, btw_bedrag=2.72, btw_soort='nl21', startdatum='2026-09-07', laatste_datum='2026-10-07', variabel=1,
        notitie='$14,99 + $3,15 btw per maand, in euro wisselt het met de koers.', key='artlist')

    buitenlands_note = ('Buitenlandse leverancier met btw: controleer op de factuur of dit Nederlandse btw is (dan aftrekbaar), '
                        'of vul je btw-nummer in zodat de btw wordt verlegd.')
    for r in range(4, 21):
        naam, datum, soort_ab, bedrag, btwtype = (exp['A%d' % r].value, exp['B%d' % r].value, exp['C%d' % r].value,
                                                   exp['D%d' % r].value, exp['F%d' % r].value)
        if not naam or not datum or bedrag is None or str(naam).startswith('START'):
            continue
        d = iso(datum)
        bedrag = float(bedrag)
        notitie = exp['K%d' % r].value or ''
        kw = dict(soort='kosten', datum=d, omschrijving='%s %s' % (naam, soort_ab or ''), relatie=naam, toewijzing='ik', betaald=1, rekening='wise',
                  categorie='Software & abonnementen', notitie=notitie)
        if naam == 'Google':
            kw.update(btw_soort='eu_verlegd', bedrag_excl=bedrag, btw_bedrag=0)
            if soort_ab == 'Standard':
                kw.update(categorie='Hosting & domeinen', abonnement_id=ab['google_j'])
            elif soort_ab == 'Monthly':
                kw.update(abonnement_id=ab['google_m'])
        elif naam == 'Freepik':
            btw = round(bedrag * 21 / 121, 2)
            kw.update(btw_soort='nl21', bedrag_excl=round(bedrag - btw, 2), btw_bedrag=btw, controleren=1,
                      notitie=(notitie + ' ' + buitenlands_note).strip())
            if soort_ab != 'One time':
                kw['abonnement_id'] = ab['freepik']
        elif naam == 'Artlist':
            if d == '2026-09-07':
                btw = 2.72
                notitie = 'Factuur: $14,99 + $3,15 btw = $18,14; afgeschreven €15,69 (+ €0,07 Wise-kosten). In je sheet stond €3,30 btw; omgerekend is dat circa €2,72.'
            else:
                btw = 2.81
                notitie = 'Zelfde dollarbedrag als september, afgeschreven €16,20 (+ €0,08 Wise-kosten). Btw omgerekend circa €2,81 (sheet: €3,30).'
            kw.update(btw_soort='nl21', bedrag_excl=round(bedrag - btw, 2), btw_bedrag=btw, controleren=1, abonnement_id=ab['artlist'],
                      notitie=notitie + ' ' + buitenlands_note)
        else:
            kw.update(btw_soort='geen', bedrag_excl=bedrag, btw_bedrag=0)
        boeking(**kw)

    # ---- stortingen en privé-uitgaven op de Wise-kaart
    for r in range(4, 12):
        if wise['A%d' % r].value and wise['C%d' % r].value == 'Added':
            fee = wise['F%d' % r].value
            boeking(soort='storting', datum=iso(wise['B%d' % r].value), omschrijving='Storting Wise-kaart', relatie='Spaarrekening (privé)',
                    bedrag_excl=wise['D%d' % r].value, rekening='wise', categorie='Privé-storting',
                    notitie='Wise-kosten voor deze storting: € %s' % fee if fee else '')
    totaal_oud = float(wise['N3'].value)
    boeking(soort='prive_uitgave', datum='2026-09-28', omschrijving='Privé-uitgaven Wise t/m 28-09-2026 (opgeteld)', bedrag_excl=totaal_oud,
            rekening='wise', categorie='Privé', notitie='Samengevat uit Finance.xlsx; losse bedragen waren niet gedateerd.')
    n = 0
    for r in range(4, 60):
        v = wise['N%d' % r].value
        if isinstance(v, (int, float)):
            n += 1
            boeking(soort='prive_uitgave', datum='2026-10-12', omschrijving='Privé-uitgave Wise (datum onbekend, na 28-09)', bedrag_excl=float(v),
                    rekening='wise', categorie='Privé', notitie='Uit de lijst "Prive spent" in Finance.xlsx.')

    c.commit()
    c.close()
    S_cfg = {'wise_startsaldo': str(wise['L7'].value)}
    with closing(S.db()) as cc:
        for k, v in S_cfg.items():
            cc.execute('INSERT OR REPLACE INTO instellingen(sleutel,waarde) VALUES(?,?)', (k, v))
        cc.commit()
    print('Import klaar: %d privé-uitgaven na 28-09, %d boekingen totaal' % (n, S.q('SELECT COUNT(*) n FROM boekingen', one=True)['n']))


if __name__ == '__main__':
    if len(sys.argv) < 2:
        sys.exit(__doc__)
    main(sys.argv[1], '--reset' in sys.argv)
